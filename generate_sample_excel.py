import pandas as pd
import numpy as np

# Daftar 35 Kecamatan Pandeglang
kecamatan_list = [
    "SUMUR", "CIMANGGU", "CIBALIUNG", "CIBITUNG", "CIKEUSIK", "CIGEULIS", "PANIMBANG", "SOBANG", 
    "MUNJUL", "ANGSANA", "SINDANGRESMI", "PICUNG", "BOJONG", "SAKETI", "CISATA", "PAGELARAN", 
    "PATIA", "SUKARESMI", "LABUAN", "CARITA", "JIPUT", "CIKEDAL", "MENES", "PULOSARI", 
    "MANDALAWANGI", "CIMANUK", "CIPEUCANG", "BANJAR", "KADUHEJO", "PANDEGLANG", "MAJASARI", 
    "KARANG TANJUNG", "CADASARI", "KORONCONG", "MEKARJAYA"
]

n_kec = len(kecamatan_list)

# Set seed untuk data deterministik yang realistis
np.random.seed(42)

# --- BAB 1: GEOGRAFI DAN IKLIM ---
bab1_data = []
for i, name in enumerate(kecamatan_list):
    luas = round(float(np.random.uniform(15.0, 320.0)), 2)
    desa_luas_luas = round(luas * float(np.random.uniform(0.1, 0.25)), 2)
    desa_luas_nama = f"DESA {name[:3]} JAYA"
    pegunungan = int(np.random.choice([0, 10, 25, 45, 60, 75, 80, 90]))
    # Kecamatan pesisir (misal Sumur, Labuan, Carita, Panimbang, Cikeusik, Cibitung, Cigeulis)
    is_pesisir = name in ["SUMUR", "LABUAN", "CARITA", "PANIMBANG", "CIKEUSIK", "CIBITUNG", "CIGEULIS"]
    pesisir = int(np.random.uniform(30, 100)) if is_pesisir else 0
    
    desa_jauh_nama = f"DESA {name[:3]} BARAT"
    desa_jauh_jarak = round(float(np.random.uniform(3.0, 35.0)), 1)
    
    desa_tinggi_nama = f"DESA {name[:3]} HILIR"
    desa_tinggi_val = int(np.random.uniform(10, 700)) if pegunungan > 40 else int(np.random.uniform(5, 120))
    
    bab1_data.append({
        "No": i + 1,
        "Kecamatan": name,
        "Luas_Kecamatan": luas,
        "Desa_Terluas_Nama": desa_luas_nama,
        "Desa_Terluas_Luas": desa_luas_luas,
        "Persen_Pegunungan": pegunungan,
        "Persen_Pesisir": pesisir,
        "Desa_Terjauh_Nama": desa_jauh_nama,
        "Desa_Terjauh_Jarak": desa_jauh_jarak,
        "Desa_Tertinggi_Nama": desa_tinggi_nama,
        "Desa_Tertinggi_Tinggi": desa_tinggi_val
    })

# --- BAB 2: PEMERINTAHAN ---
bab2_data = []
for i, name in enumerate(kecamatan_list):
    pns_total = int(np.random.randint(10, 45))
    pns_laki = int(np.random.randint(40, 75))
    pns_perempuan = 100 - pns_laki
    
    p_sd_smp = round(float(np.random.uniform(0, 10)), 1)
    p_sma = round(float(np.random.uniform(10, 30)), 1)
    p_diploma = round(float(np.random.uniform(5, 20)), 1)
    p_sarjana = round(100.0 - (p_sd_smp + p_sma + p_diploma), 1)
    
    g_iv = int(np.random.randint(5, 20))
    g_iii = int(np.random.randint(50, 75))
    g_ii = int(np.random.randint(10, 25))
    g_i = 100 - (g_iv + g_iii + g_ii)
    
    bab2_data.append({
        "No": i + 1,
        "Kecamatan": name,
        "PNS_Total": pns_total,
        "PNS_Laki": pns_laki,
        "PNS_Perempuan": pns_perempuan,
        "Pendidikan_SD_SMP": p_sd_smp,
        "Pendidikan_SMA": p_sma,
        "Pendidikan_Diploma": p_diploma,
        "Pendidikan_Sarjana": p_sarjana,
        "Golongan_I": g_i,
        "Golongan_II": g_ii,
        "Golongan_III": g_iii,
        "Golongan_IV": g_iv
    })

# --- BAB 3: PENDUDUK ---
bab3_data = []
for i, name in enumerate(kecamatan_list):
    luas = bab1_data[i]["Luas_Kecamatan"]
    penduduk = int(np.random.randint(15000, 65000))
    kepadatan = round(penduduk / luas, 1)
    
    p_laki = round(float(np.random.uniform(49.0, 52.5)), 1)
    p_perempuan = round(100.0 - p_laki, 1)
    sex_ratio = round((p_laki / p_perempuan) * 100, 1)
    
    p_muda = round(float(np.random.uniform(22.0, 29.0)), 1)
    p_produktif = round(float(np.random.uniform(62.0, 68.0)), 1)
    p_tua = round(100.0 - (p_muda + p_produktif), 1)
    beban = round(((p_muda + p_tua) / p_produktif) * 100, 1)
    
    desa_max_nama = f"DESA {name[:3]} MAJU"
    desa_max_val = int(penduduk * float(np.random.uniform(0.12, 0.22)))
    
    desa_padat_nama = f"DESA {name[:3]} INDAH"
    desa_padat_val = round(kepadatan * float(np.random.uniform(1.5, 3.0)), 1)
    
    bab3_data.append({
        "No": i + 1,
        "Kecamatan": name,
        "Penduduk_Total": penduduk,
        "Kepadatan": kepadatan,
        "Persen_Laki": p_laki,
        "Persen_Perempuan": p_perempuan,
        "Sex_Ratio": sex_ratio,
        "Persen_Usia_Muda": p_muda,
        "Persen_Usia_Produktif": p_produktif,
        "Persen_Usia_Lanjut": p_tua,
        "Beban_Tanggungan": beban,
        "Desa_Penduduk_Terbesar_Nama": desa_max_nama,
        "Desa_Penduduk_Terbesar_Val": desa_max_val,
        "Desa_Kepadatan_Tertinggi_Nama": desa_padat_nama,
        "Desa_Kepadatan_Tertinggi_Val": desa_padat_val
    })

# --- BAB 4: SOSIAL DAN KESEJAHTERAAN RAKYAT ---
bab4_data = []
for i, name in enumerate(kecamatan_list):
    f_tk = int(np.random.randint(2, 8))
    f_sd = int(np.random.randint(10, 28))
    f_smp = int(np.random.randint(2, 7))
    f_sma = int(np.random.randint(1, 5))
    
    s_tk = f_tk * int(np.random.randint(30, 60))
    s_sd = f_sd * int(np.random.randint(110, 160))
    s_smp = f_smp * int(np.random.randint(80, 130))
    s_sma = f_sma * int(np.random.randint(70, 120))
    
    g_kurang = int(np.random.randint(40, 150))
    g_buruk = int(np.random.randint(5, 45))
    g_kurus = int(np.random.randint(30, 140))
    g_stunting = int(np.random.randint(8, 35))
    
    jalan = int(np.random.choice([40, 50, 60, 75, 80, 90, 100]))
    
    bab4_data.append({
        "No": i + 1,
        "Kecamatan": name,
        "Fasilitas_TK": f_tk,
        "Fasilitas_SD": f_sd,
        "Fasilitas_SMP": f_smp,
        "Fasilitas_SMA": f_sma,
        "Siswa_TK": s_tk,
        "Siswa_SD": s_sd,
        "Siswa_SMP": s_smp,
        "Siswa_SMA": s_sma,
        "Gizi_Kurang": g_kurang,
        "Gizi_Buruk": g_buruk,
        "Gizi_Kurus": g_kurus,
        "Gizi_Stunting": g_stunting,
        "Penerangan_Jalan": jalan
    })

# --- BAB 5: PERTANIAN ---
bab5_data = []
for i, name in enumerate(kecamatan_list):
    panen_k = int(np.random.randint(0, 15))
    panen_r = int(np.random.randint(0, 20))
    
    prod_k = panen_k * int(np.random.randint(80, 120)) if panen_k > 0 else 0
    prod_r = panen_r * int(np.random.randint(70, 110)) if panen_r > 0 else 0
    
    prod_lengkeng = int(np.random.randint(50, 1800))
    prod_duku = int(np.random.randint(100, 6000))
    
    bab5_data.append({
        "No": i + 1,
        "Kecamatan": name,
        "Panen_Cabai_Keriting": panen_k,
        "Panen_Cabai_Rawit": panen_r,
        "Produksi_Cabai_Keriting": prod_k,
        "Produksi_Cabai_Rawit": prod_r,
        "Produksi_Lengkeng": prod_lengkeng,
        "Produksi_Duku": prod_duku
    })

# --- BAB 6: PARIWISATA, TRANSPORTASI, DAN KOMUNIKASI ---
bab6_data = []
for i, name in enumerate(kecamatan_list):
    wisata = int(np.random.randint(0, 8))
    bts = int(np.random.randint(1, 12))
    sinyal = int(np.random.choice([50, 60, 70, 80, 90, 100]))
    angkot = int(np.random.randint(0, 450))
    jalan_kond = str(np.random.choice(["BAIK", "SEDANG", "RUSAK"], p=[0.4, 0.45, 0.15]))
    
    bab6_data.append({
        "No": i + 1,
        "Kecamatan": name,
        "Wisata_Objek": wisata,
        "Menara_BTS": bts,
        "Persen_Sinyal_Kuat": sinyal,
        "Jumlah_Angkot": angkot,
        "Kondisi_Jalan": jalan_kond
    })

# --- BAB 7: PERBANKAN, KOPERASI, DAN PERDAGANGAN ---
bab7_data = []
for i, name in enumerate(kecamatan_list):
    koperasi = int(np.random.randint(0, 12))
    bank = int(np.random.randint(0, 5))
    toko = int(np.random.randint(30, 350))
    minimarket = int(np.random.randint(0, 15))
    industri = int(np.random.randint(5, 120))
    
    bab7_data.append({
        "No": i + 1,
        "Kecamatan": name,
        "Koperasi_Aktif": koperasi,
        "Jumlah_Bank": bank,
        "Pasar_Toko": toko,
        "Minimarket": minimarket,
        "Industri_Mikro": industri
    })

# --- SIMPAN KE EXCEL DENGAN 7 SHEET (WRITER) ---
import openpyxl
from openpyxl.utils import get_column_letter

output_file = "sample_data_kecamatan.xlsx"
bab_titles = {
    1: "Geografi dan Iklim",
    2: "Pemerintahan",
    3: "Penduduk",
    4: "Sosial dan Kesejahteraan Rakyat",
    5: "Pertanian",
    6: "Pariwisata, Transportasi, dan Komunikasi",
    7: "Perbankan, Koperasi, dan Perdagangan"
}

with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
    for b in range(1, 8):
        sheet_name = f"Bab {b}"
        df = pd.DataFrame(eval(f"bab{b}_data"))
        
        # Tulis dataframe mulai dari baris 5 (startrow=4)
        df.to_excel(writer, sheet_name=sheet_name, index=False, startrow=4)
        
        # Tambahkan judul dan petunjuk di baris 1-3
        ws = writer.sheets[sheet_name]
        ws.cell(row=1, column=1, value="TEMPLATE DATA MASUKAN INFOGRAFIS KCDA 2024 - BPS KABUPATEN PANDEGLANG")
        ws.cell(row=2, column=1, value=f"{sheet_name.upper()}: {bab_titles[b].upper()}")
        ws.cell(row=3, column=1, value="Petunjuk: Isi data tiap kecamatan di bawah ini. Jangan mengubah baris header (Baris 5)!")
        
        # Autofit lebar kolom
        for col in ws.columns:
            # Lewati baris 1-3 yang sangat panjang saat menghitung lebar
            cell_lens = []
            for cell in col:
                if cell.row > 4 and cell.value is not None:
                    cell_lens.append(len(str(cell.value)))
            max_len = max(cell_lens) if cell_lens else 10
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

print(f"File Excel multi-sheet berhasil dibuat di {output_file}")
